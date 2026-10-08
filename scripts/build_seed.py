import json, re, unicodedata, sys
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],read_only=True,data_only=True)
def norm(s):
    s=unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().lower()
    s=re.sub(r'\(.*?\)','',s); return re.sub(r'\s+',' ',s).strip()
def clean(v):
    if isinstance(v,str):
        v=re.sub(r'[ \t]+',' ',v).replace('/ ','/').strip(); return v or None
    return v
ALIAS={'especialistas de dados':'especialista de dados','especialista em inteligencia artificial':'especialista de ia',
 'coordenador de compliance':'coordenador de qualidade de midia'}
def key(s): k=norm(s); return ALIAS.get(k,k)
HEAD={'CARGO','Cargo de Origem'}
cargos={}
for r in wb['Descrição Cargos'].iter_rows(values_only=True):
    r=[clean(c) for c in (list(r)+[None]*11)[1:11]]
    if not r[0] or r[0] in HEAD: continue
    nome,desc,foco,compl,aut,papel,req,time,lider,area=r
    if req and 'Secretária Executiva Sr' in req: req=req.split('Secretária Executiva Sr')[0].strip()
    cargos[key(nome)]=dict(nome=nome,descricao=desc,foco=foco,complexidade=compl,autonomia=aut,papel=papel,requisitos=req,time=time,lider=lider,area=area)
for r in wb['Cargos e Salários'].iter_rows(values_only=True):
    r=[clean(c) for c in (list(r)+[None]*9)[1:9]]
    if not r[2] or r[2] in HEAD: continue
    area,time,nome,pos,nivel,papel,piso,teto=r
    c=cargos.setdefault(key(nome),dict(nome=nome))
    c.update(position=pos,nivel_raw=nivel,piso=piso,teto=teto)
    if not c.get('area'): c['area']=area
    if not c.get('time'): c['time']=time
    if not c.get('papel'): c['papel']=papel
def nivel(c):
    n=c['nome']; raw=c.get('nivel_raw') or ''
    if n.startswith('Coordenador') or raw=='Coordenador': return 'Coordenador'
    if n.startswith('Especialista') or raw=='Especialista': return 'Especialista'
    for pat,lv in [(r'\b(Jr|Júnior)\b','Júnior'),(r'\b(Pl|Pleno)\b','Pleno'),(r'\b(Sr|Sênior)\b','Sênior')]:
        if re.search(pat,n) or re.search(pat,raw): return lv
    return 'Júnior' if n=='Secretária Executiva' else 'Especialista'
FAM={'Analista de CRO':'CRO','Analista de Dados':'Dados','Especialista de Dados':'Dados','Especialistas de Dados':'Dados',
 'Analista de Growth':'Growth','Analista de Infraestrutura':'Infraestrutura','Especialista de Infraestrutura':'Infraestrutura',
 'Analista de Produto':'Produto','Especialista de Produto':'Produto','Analista de Qualidade de Mídia':'Qualidade de Mídia',
 'Coordenador de Compliance':'Qualidade de Mídia','Coordenador de Qualidade de Mídia':'Qualidade de Mídia',
 'Analista de Redes Sociais':'Redes Sociais','Especialista de Redes Sociais':'Redes Sociais','Analista de RH':'RH','Coordenador de RH':'RH',
 'Analista de Sucesso do Cliente':'Sucesso do Cliente','Coordenador de Sucesso do Cliente':'Sucesso do Cliente',
 'Analista de Tráfego Pago':'Tráfego Pago','Especialista de Tráfego Pago':'Tráfego Pago','Coordenador de Performance':'Tráfego Pago',
 'Secretária Executiva':'Secretariado Executivo','Coordenadora Executiva':'Secretariado Executivo',
 'Analista Financeiro':'Financeiro','Coordenador Financeiro':'Financeiro','Coordenador de Conteúdo':'Redação','Redator(a)':'Redação',
 'Especialista de Redação':'Redação','Coordenador de Tecnologia':'Desenvolvimento de Software','Desenvolvedor de Software':'Desenvolvimento de Software',
 'Especialista de Software':'Desenvolvimento de Software','Desenvolvedor Web':'Desenvolvimento Web','Especialista em Desenvolvimento Web':'Desenvolvimento Web',
 'Designer UX/UI':'Design UX/UI','Especialista de Design':'Design','Especialista de Projetos':'Projetos',
 'Especialista em Inteligência Artificial (IA)':'Inteligência Artificial'}
def familia(n):
    m=[p for p in FAM if n.startswith(p)]
    return FAM[max(m,key=len)] if m else n
out=[]
for c in cargos.values():
    c['nivel']=nivel(c); c['familia']=familia(c['nome']); c.pop('nivel_raw',None)
    for f in ['position','piso','teto','descricao','foco','complexidade','autonomia','papel','requisitos','lider','area','time']: c.setdefault(f,None)
    for f in ['piso','teto']:
        if c[f] is not None: c[f]=int(c[f])
    out.append(c)
names={key(c['nome']):c['nome'] for c in out}
mob=[]; missing=set()
for r in wb['Trilhas Horizontais'].iter_rows(values_only=True):
    o,d,t,obs=[clean(x) for x in (list(r)+[None]*5)[1:5]]
    if not o or not d or o in HEAD: continue
    ko,kd=key(o),key(d)
    for kk,raw in ((ko,o),(kd,d)):
        if kk not in names: missing.add(raw)
    if ko in names and kd in names: mob.append(dict(origem=names[ko],destino=names[kd],tipo=t,observacao=obs))
print(len(out),'cargos',len(mob),'mobilidades; não encontrados:',missing)
for c in sorted(out,key=lambda c:(c['familia'],c['nome'])): print(c['familia'],'|',c['nome'],'|',c['nivel'],'|',c['area'],'|',c['time'],c['piso'],c['teto'],c['lider'])
json.dump(dict(cargos=out,mobilidades=mob),open('data/seed.json','w'),ensure_ascii=False,indent=1)
